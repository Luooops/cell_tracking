"""Build the two GT audit workbooks from recursive evaluation outputs.

The existing four audit workbooks supply the layout and the old-GT baseline.
CSV records are checked against the XML named by each evaluation summary.
"""
import argparse
import csv
import hashlib
import json
import math
import os
import re
import xml.etree.ElementTree as ET
from pathlib import Path
import shutil
import subprocess
import tempfile
from collections import defaultdict
from datetime import datetime

import numpy as np
import openpyxl  # Read templates/baselines only; Node writes the workbooks.

import zipfile

HERE = Path(__file__).resolve().parent
NAME = re.compile(r'(r\d+c\d+)(f\d+)(p\d+)-(ch\d+)t(\d+)\.[^.]+$', re.I)
TEMPLATES = {
    'dup': 'GT_new_重复ID汇总_20260923.xlsx',
    'disp': 'GT_new_单帧位移异常分析_20260923.xlsx',
    'old_dup': 'GT_XML重复ID汇总.xlsx',
    'old_disp': 'GT_单帧位移异常分析_20260915.xlsx',
}


def read_gt(path):
    # Do not import the evaluation engine: exporting needs no SciPy/tracker runtime.
    if not zipfile.is_zipfile(path):
        return ET.parse(path).getroot()
    with zipfile.ZipFile(path) as archive:
        members = [m for m in archive.infolist() if not m.is_dir() and m.filename.lower().endswith('.xml')]
        if len(members) != 1:
            raise ValueError(f'Expected exactly one XML inside {path}; got {len(members)}')
        with archive.open(members[0]) as stream:
            return ET.parse(stream).getroot()


def centroid(points):
    """Same area-centroid definition as evaluate.py, without engine dependencies."""
    x, y = points.T
    cross = x*np.roll(y,-1)-np.roll(x,-1)*y
    area2 = cross.sum()
    if abs(area2) < 1e-8:
        return points.mean(axis=0)
    return np.array([((x+np.roll(x,-1))*cross).sum(),
                     ((y+np.roll(y,-1))*cross).sum()])/(3*area2)


def read_books(directory):
    books = {}
    for key, filename in TEMPLATES.items():
        with (directory / filename).open('rb') as stream:
            book = openpyxl.load_workbook(stream, data_only=True)
            books[key] = {s.title: [list(row) for row in s.values] for s in book}
            book.close()
    return books


def discover(outputs, gt_root=None):
    """Fail on incomplete or overlapping runs rather than silently double-count."""
    jobs, seen = [], set()
    for manifest in sorted(outputs.rglob('summary.json')):
        report = json.loads(manifest.read_text(encoding='utf-8-sig'))
        if 'gt' not in report or 'dataset' not in report:
            continue  # Other output types are not evaluation runs.
        csv_path = manifest.parent / 'position_matches.csv'
        if not csv_path.is_file():
            raise ValueError(f'Missing position_matches.csv: {manifest.parent}')
        gt = Path(report['gt'])
        if gt_root:
            candidates = list((gt_root / report['dataset']).rglob(gt.name))
            if len(candidates) != 1:
                raise ValueError(f'Expected one relocated GT for {manifest}, got {len(candidates)}')
            gt = candidates[0]
        if gt.is_dir():
            candidates = [p for p in gt.iterdir() if p.suffix.lower() in {'.xml', '.zip'}
                          and (p.suffix.lower() == '.xml' or has_xml(p))]
            if len(candidates) != 1:
                raise ValueError(f'Expected one GT XML/ZIP in {gt}')
            gt = candidates[0]
        key = (report['dataset'], gt.resolve())
        if key in seen:
            raise ValueError(f'Multiple evaluation runs reference {gt}; select one outputs tree.')
        seen.add(key)
        if not gt.is_file():
            raise FileNotFoundError(f'GT missing: {gt}. Use --gt-root if relocated.')
        jobs.append((manifest, report, gt, csv_path))
    if not jobs:
        raise ValueError(f'No evaluation summary.json files in {outputs}')
    return jobs


def has_xml(path):
    with zipfile.ZipFile(path) as archive:
        return any(n.lower().endswith('.xml') for n in archive.namelist())


def check_snapshot(csv_path, observations):
    expected = {(o['name'], o['index']): o for o in observations if o['label'] == 'cell'}
    seen = set()
    with csv_path.open(encoding='utf-8-sig', newline='') as stream:
        for row in csv.DictReader(stream):
            key = (row['image_name'], int(row['gt_instance_id']))
            o = expected.get(key)
            if o is None or key in seen:
                raise ValueError(f'CSV/XML object mismatch: {csv_path}: {key}')
            if (row['gt_track_id'].strip() != o['tid'] or
                    row['xml_frame_id'] != o['xml_id'] or int(row['time_index']) != o['t'] or
                    not math.isclose(float(row['gt_x']), o['x'], abs_tol=1e-6, rel_tol=0) or
                    not math.isclose(float(row['gt_y']), o['y'], abs_tol=1e-6, rel_tol=0)):
                raise ValueError(f'CSV and XML are different GT snapshots: {csv_path}: {key}')
            seen.add(key)
    if seen != expected.keys():
        raise ValueError(f'CSV misses {len(expected.keys() - seen)} XML objects: {csv_path}')


def transitions(times):
    """Do not bridge gaps or choose arbitrarily between duplicate endpoints."""
    ordered = sorted(times)
    for a, b in zip(ordered, ordered[1:]):
        left, right = times[a], times[b]
        reason = ('缺帧/跨帧' if b-a != 1 else
                  '端点同帧重复ID' if len(left) != 1 or len(right) != 1 else None)
        distance = None if reason else math.hypot(right[0]['x']-left[0]['x'], right[0]['y']-left[0]['y'])
        yield a, b, left, right, reason, distance


def quantile(values, q):
    return float(np.quantile(values, q)) if values else None


def ratio(a, b):
    return a/b if b else None


def analyze(jobs):
    groups = defaultdict(lambda: defaultdict(list))
    wells, sources, dup_summary, dup_events, dup_objects = {}, [], [], [], []
    seen_frames = set()
    for manifest, report, path, csv_path in jobs:
        root = read_gt(path)
        images = root.findall('image')
        if not images:
            raise ValueError(f'No per-image annotations: {path}')
        sid, dataset = f'S{len(sources)+1:02d}', report['dataset']
        member = ''
        if zipfile.is_zipfile(path):
            with zipfile.ZipFile(path) as archive:
                member = next(n for n in archive.namelist() if n.lower().endswith('.xml'))
        observations, conflicts, duplicate_frames, per_track = [], set(), set(), defaultdict(list)
        file_wells, event_start, object_start = set(), len(dup_events), len(dup_objects)
        missing = cell_missing = other = 0
        for im in images:
            name = im.attrib['name'].replace('\\', '/').split('/')[-1]
            match = NAME.fullmatch(name)
            if not match:
                raise ValueError(f'Unsupported image name: {name}')
            well, field, plane, channel, time = match.groups()
            t, seq = int(time), f'{field}/{plane}/{channel}'
            frame_key = dataset, well, seq, t
            if frame_key in seen_frames:
                raise ValueError(f'Overlapping GT frames from multiple runs: {frame_key}')
            seen_frames.add(frame_key)
            file_wells.add(well)
            w = wells.setdefault((dataset, well), dict(frames=0, objects=0, missing=0, degenerate=0, seqs=set(), sources=set()))
            w['frames'] += 1; w['seqs'].add(seq); w['sources'].add(sid)
            ids = defaultdict(list)
            for index, polygon in enumerate(im.findall('polygon'), 1):
                attrs = [(a.text or '').strip() for a in polygon.findall('attribute') if a.get('name') == 'track_id']
                if len(attrs) > 1:
                    raise ValueError(f'Multiple track_id attributes: {name} #{index}')
                tid = attrs[0] if attrs else ''
                points = np.array([list(map(float, p.split(','))) for p in polygon.attrib['points'].split(';')])
                if points.ndim != 2 or points.shape[1] != 2 or len(points) < 3 or not np.isfinite(points).all():
                    raise ValueError(f'Invalid polygon: {name} #{index}')
                x, y = map(float, centroid(points))
                width, height = int(im.attrib['width']), int(im.attrib['height'])
                if width <= 0 or height <= 0:
                    raise ValueError(f'Invalid image dimensions: {name}')
                box = [float(points[:,0].min()), float(points[:,1].min()), float(points[:,0].max()), float(points[:,1].max())]
                region = ['上','中','下'][min(2,max(0,int(y*3/height)))]+['左','中','右'][min(2,max(0,int(x*3/width)))]
                o = dict(name=name, index=index, tid=tid, x=x, y=y, xml_id=im.get('id',''), t=t,
                         source=sid, width=width, height=height, label=polygon.get('label',''), box=box, region=region)
                observations.append(o)
                if tid:
                    ids[tid].append(o); per_track[(well,seq,tid)].append(o)
                else:
                    missing += 1
                if o['label'] != 'cell':
                    other += 1; continue
                w['objects'] += 1
                px, py = points.T
                w['degenerate'] += int(abs((px*np.roll(py,-1)-np.roll(px,-1)*py).sum()) < 1e-8)
                if not tid:
                    cell_missing += 1; w['missing'] += 1; continue
                groups[(dataset,well,seq,tid)][t].append(o)
            for tid, objects in sorted(ids.items()):
                if len(objects) < 2:
                    continue
                eid = f'E{len(dup_events)+1:05d}'
                conflicts.add((well,seq,tid)); duplicate_frames.add((well,seq,t))
                description = '; '.join(f"#{o['index']} ({o['x']:.1f}, {o['y']:.1f}) {o['region']}" for o in objects)
                dup_events.append([eid,dataset,well,f'{well}/{seq}',tid,t,im.get('id',''),name,len(objects),description,str(path),member])
                for o in objects:
                    dup_objects.append([eid,dataset,f'{well}/{seq}',tid,t,name,o['index'],o['label'],o['x'],o['y'],o['region'],*o['box'],width,height,str(path),member,im.get('id','')])
        check_snapshot(csv_path, observations)
        if report['gt_frames'] != len(images) or report['gt_objects'] != len(observations)-other:
            raise ValueError(f'Summary and XML counts differ: {manifest}')
        dup_summary.append([dataset,','.join(sorted(file_wells)),len(images),len(observations),missing,len(conflicts),len(duplicate_frames),len(dup_events)-event_start,len(dup_objects)-object_start,sum(len(per_track[k]) for k in conflicts),'已扫描',str(path),member])
        sources.append([sid,dataset,','.join(sorted(file_wells)),str(path),member,hashlib.sha256(path.read_bytes()).hexdigest(),len(images),len(observations)-other,cell_missing,other])
        print(f'{sid}: {path.name} ({len(images)} frames; CSV/XML checked)', flush=True)

    pairs, info, exclusions = [], {}, []
    for key, times in sorted(groups.items()):
        dataset, well, seq, tid = key
        v = dict(start=min(times), end=max(times), observed=len(times), duplicates=sum(len(x)>1 for x in times.values()), pairs=[], gaps=0, ambiguous=0)
        info[key] = v
        for a,b,aa,bb,reason,distance in transitions(times):
            p,r = aa[0],bb[0]
            if reason:
                v['gaps' if b-a != 1 else 'ambiguous'] += 1
                exclusions.append([well,tid,a,b,b-a,reason,len(aa),len(bb),','.join(str(o['index']) for o in aa),','.join(str(o['index']) for o in bb),p['name'],r['name'],seq,dataset,p['source'],r['source']])
            else:
                item = dict(key=key,a=a,b=b,p=p,r=r,d=distance,conflict=bool(v['duplicates']))
                v['pairs'].append(item); pairs.append(item)
    summary, events, tracks = [], [], []
    for wk,w in sorted(wells.items()):
        local = [p for p in pairs if p['key'][:2] == wk]
        values = [p['d'] for p in local]
        median = quantile(values,.5); threshold = median*4 if median is not None else None
        p999 = quantile(values,.999)
        iv = [v for k,v in info.items() if k[:2] == wk]
        flagged = set()
        for p in local:
            p['flag'] = p['d'] > threshold
            if not p['flag']: continue
            flagged.add(p['key'])
            dataset,well,seq,tid = p['key']; a,b = p['p'],p['r']
            events.append([well,tid,p['a'],p['b'],p['d'],threshold,ratio(p['d'],threshold),ratio(p['d'],median),a['x'],a['y'],b['x'],b['y'],b['x']-a['x'],b['y']-a['y'],ratio(p['d'],a['width']),p999,ratio(p['d'],p999),'是' if p['conflict'] else '否',a['xml_id'],b['xml_id'],a['index'],b['index'],a['name'],b['name'],seq,dataset,a['source'],b['source']])
        clean = [p for p in local if not p['conflict']]
        ct = quantile([p['d'] for p in clean],.5); ct = ct*4 if ct is not None else None
        clean_bad = {p['key'] for p in clean if p['d'] > ct}
        eligible = sum(bool(v['pairs']) for v in iv)
        summary.append([*wk,len(w['seqs']),w['frames'],w['objects'],w['missing'],len(iv),eligible,sum(v['duplicates']>0 for v in iv),len(values),median,threshold,quantile(values,.99),p999,quantile([d for d in values if d<=threshold],.999),max(values) if values else None,sum(p['flag'] for p in local),len(flagged),ratio(len(flagged),len(iv)),ratio(len(flagged),eligible),sum(v['gaps'] for v in iv),sum(v['ambiguous'] for v in iv),len(clean_bad),sum(bool(v['pairs']) and not v['duplicates'] for v in iv),ct,w['degenerate'],','.join(sorted(w['sources']))])
    for key,v in info.items():
        dataset,well,seq,tid = key; ps = v['pairs']; bad = sum(p['flag'] for p in ps)
        m = max(ps,key=lambda p:p['d']) if ps else None
        tracks.append([well,tid,'异常' if bad else ('无可评估帧对' if not ps else '未检出'),bad,m['d'] if m else None,m['a'] if m else None,m['b'] if m else None,v['start'],v['end'],v['observed'],v['duplicates'],len(ps),v['gaps'],v['ambiguous'],seq,dataset])
    events.sort(key=lambda r:-r[4]); tracks.sort(key=lambda r:(-r[3],-(r[4] or 0)))
    assert len(events) == sum(r[16] for r in summary)
    assert sum(r[17] for r in summary) == sum(r[3]>0 for r in tracks)
    assert len(dup_events) == sum(r[7] for r in dup_summary)
    assert len(dup_objects) == sum(r[8] for r in dup_summary) == sum(r[8] for r in dup_events)
    return dict(summary=summary,events=events,tracks=tracks,exclusions=exclusions,sources=sources,
                dup_summary=dup_summary,dup_events=dup_events,dup_objects=dup_objects), pairs


def compare(data, pairs, books):
    old_ids = {(r[0],r[1]):r for r in books['old_dup']['文件汇总'][1:] if r[0]}
    source_rows = books['old_disp']['口径与来源']
    start = next(i for i,r in enumerate(source_rows) if r[0]=='来源编号')+1
    old_sources = {r[0]:r for r in source_rows[start:] if r[0]}
    old_main = {}
    rows = books['old_disp']['Well汇总']; start = next(i for i,r in enumerate(rows) if r[0]=='Well' and r[1]=='全部 track')+1
    for row in rows[start:]:
        if not row[0] or row[0]=='合计': break
        source = old_sources[row[10].split(',')[0]]
        old_main[(source[1],row[0])] = row
    result = []
    for row in data['summary']:
        key = tuple(row[:2]); old = old_main.get(key); duplicate = old_ids.get(key)
        new_dup = sum(r[7] for r in data['dup_summary'] if tuple(r[:2])==key)
        new_conflicts = row[8]
        if old is None or duplicate is None:
            result.append([row[1],None,new_dup,None,new_conflicts,None,row[16],None,row[17],None,row[19],None,row[11],None,row[15],None,None,row[0]])
            continue
        previous_gate = [p for p in pairs if p['key'][:2]==key and p['d']>old[8]]
        result.append([row[1],duplicate[7],new_dup,duplicate[5],new_conflicts,old[6],row[16],old[3],row[17],old[5],row[19],old[8],row[11],old[9],row[15],len(previous_gate),len({p['key'] for p in previous_gate}),row[0]])
    return result


def export(data, args):
    runtime = Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
    node = args.node or (runtime/'bin/node.exe' if (runtime/'bin/node.exe').is_file() else shutil.which('node'))
    modules = args.node_modules or runtime/'node_modules'
    if not node or not (modules/'@oai/artifact-tool').is_dir():
        raise ValueError('Node and @oai/artifact-tool required; use --node and --node-modules.')
    names = [f'GT_outputs_重复ID汇总_{args.date}.xlsx',f'GT_outputs_单帧位移异常分析_{args.date}.xlsx']
    args.output_dir.mkdir(parents=True,exist_ok=True)
    for name in names:
        if (args.output_dir/name).exists() and not args.overwrite:
            raise FileExistsError(f'{args.output_dir/name} already exists; use --overwrite.')
    with tempfile.TemporaryDirectory(prefix='gt-reports-') as temporary:
        work = Path(temporary)
        (work/'export_gt_reports.mjs').write_text(WORKBOOK_RENDERER, encoding='utf-8')
        # Node resolves from this temporary package without modifying dependency dirs.
        junction = work/'node_modules'
        if os.name == 'nt':
            command = "New-Item -ItemType Junction -Path '{}' -Target '{}' | Out-Null".format(str(junction).replace("'","''"),str(modules.resolve()).replace("'","''"))
            subprocess.run(['powershell','-NoProfile','-Command',command],check=True)
        else:
            junction.symlink_to(modules.resolve(),target_is_directory=True)
        try:
            (work/'data.json').write_text(json.dumps(data,ensure_ascii=False,allow_nan=False),encoding='utf-8')
            subprocess.run([str(node),str(work/'export_gt_reports.mjs')],check=True)
            # Read exported files back before touching final destinations.
            for name,key in zip(['duplicates.xlsx','displacements.xlsx'],['dup','disp']):
                book = openpyxl.load_workbook(work/name,data_only=True)
                sheet = book['重复帧及位置' if key=='dup' else '异常跳变明细']
                expected = data['dup_events' if key=='dup' else 'events']
                if sheet.max_row != len(expected)+1:
                    raise ValueError(f'Export row count mismatch: {name}')
                for actual,wanted in zip(sheet.iter_rows(min_row=2,values_only=True),expected):
                    for a,b in zip(actual,wanted):
                        if isinstance(b,(int,float)):
                            if not isinstance(a,(int,float)) or not math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-8): raise ValueError(f'Export value mismatch: {name}')
                        elif (a or '') != (b or ''): raise ValueError(f'Export text mismatch: {name}')
                for tab in book:
                    if any(c.data_type == 'e' for row in tab for c in row):
                        raise ValueError(f'Excel error cell: {name}/{tab.title}')
                book.close()
            for source,target in zip(['duplicates.xlsx','displacements.xlsx'],names):
                shutil.copyfile(work/source,args.output_dir/target)
                print(f'Saved: {args.output_dir/target}')
        finally:
            # Remove the junction itself before TemporaryDirectory cleanup.
            if os.name == 'nt': os.rmdir(junction)
            else: junction.unlink()


# Embedded renderer is materialized only inside the auto-cleaned temporary directory.
WORKBOOK_RENDERER = r'''// Workbook layout renderer. Called by export_gt_reports.py with a checked JSON snapshot.
import fs from 'node:fs/promises';
import {fileURLToPath} from 'node:url';
import {Workbook, SpreadsheetFile} from '@oai/artifact-tool';

const base = fileURLToPath(new URL('.', import.meta.url));
const d = JSON.parse(await fs.readFile(base + 'data.json', 'utf8'));
function col(i) {
  let name = '';
  for (i++; i; i = Math.floor((i-1)/26)) name = String.fromCharCode(65+(i-1)%26)+name;
  return name;
}
function workbook(names) {
  const w = Workbook.create();
  for (const name of names) w.worksheets.add(name).showGridLines = false;
  return w;
}
function table(w, name, row, headers, rows, widths = []) {
  const s = w.worksheets.getItem(name), end = row+rows.length, last = col(headers.length-1);
  const range = s.getRange(`A${row}:${last}${end}`);
  range.values = [headers,...rows];
  range.format.font = {name:'Arial',size:10,color:'#243247'};
  range.format.rowHeight = 22;
  range.format.verticalAlignment = 'center';
  headers.forEach((_,i) => {s.getRange(`${col(i)}${row}:${col(i)}${end}`).format.columnWidth = widths[i] || 18;});
  if (rows.length) {
    const t = s.tables.add(`A${row}:${last}${end}`,true,`Table_${++table.index}`);
    t.showFilterButton = true;
  }
  s.getRange(`A${row}:${last}${row}`).format = {
    fill:'#243B5A',font:{name:'Arial',size:10,bold:true,color:'#FFFFFF'},
    rowHeight:42,wrapText:true,horizontalAlignment:'center',verticalAlignment:'center'
  };
  if (row===1) s.freezePanes.freezeRows(1);
  return s;
}
table.index = 0;
function fmt(s, columns, start, end, format) {
  if (end >= start) for (const c of columns) s.getRange(`${c}${start}:${c}${end}`).setNumberFormat(format);
}
function title(s, text, note) {
  s.getRange('A2').values = [[text]];
  s.getRange('A2').format.font = {name:'Arial',size:16,bold:true};
  s.getRange('A3').values = [[note]];
}
function totals(s, row, columns, start=6, unavailable=[]) {
  s.getRange(`A${row}`).values = [['合计']];
  for (const c of columns) {
    if (unavailable.includes(c)) s.getRange(`${c}${row}`).values = [['n.a.']];
    else s.getRange(`${c}${row}`).formulas = [[`=SUM(${c}${start}:${c}${row-1})`]];
  }
  s.getRange(`A${row}:${columns.at(-1)}${row}`).format.font = {bold:true};
}
function delta(a,b) {return a===null || b===null ? null : a-b;}
async function save(w,name) {
  w.recalculate();
  console.log((await w.inspect({kind:'match',searchTerm:'#REF!|#DIV/0!|#VALUE!|#NAME\\?|#NUM!|#SPILL!',options:{useRegex:true,maxResults:20}})).ndjson);
  for (const sheetName of Object.keys(d.books[name==='duplicates'?'dup':'disp'])) {
    const s = w.worksheets.getItem(sheetName);
    const range = s.name.includes('口径') ? 'A1:B7' : s.name.includes('对比') || s.name==='Well汇总' ? 'A1:H10' : 'A1:H7';
    const p = await w.render({sheetName:s.name,range,scale:1,format:'png'});
    await fs.writeFile(base+name+'_'+s.name+'.png',new Uint8Array(await p.arrayBuffer()));
  }
  await (await SpreadsheetFile.exportXlsx(w)).save(base+name+'.xlsx');
}

const dup = workbook(Object.keys(d.books.dup));
const dupHeader = name => d.books.dup[name][0];
table(dup,'文件汇总',1,dupHeader('文件汇总'),d.dup_summary.map((r,i)=>[...r,d.sources[i][5]]),[42,13,12,18,17,17,17,17,20,22,15,110,24,68]);
const de = table(dup,'重复帧及位置',1,dupHeader('重复帧及位置'),d.dup_events,[15,42,13,27,14,14,15,44,17,105,110,24]);
if (d.dup_events.length) de.getRange(`J2:J${d.dup_events.length+1}`).format.wrapText = true;
d.dup_events.forEach((r,i)=>{if(r[8]>2)de.getRange(`A${i+2}:L${i+2}`).format.rowHeight=40;});
const ob = table(dup,'重复对象明细',1,dupHeader('重复对象明细'),d.dup_objects,[15,42,27,14,14,44,21,15,19,19,18,18,18,18,18,18,18,110,24,15]);
fmt(ob,['I','J','L','M','N','O'],2,d.dup_objects.length+1,'0.00');
const dc = table(dup,'与旧GT对比',5,d.books.dup['与旧GT对比'][4],d.compare.map(r=>[r[0],r[1],r[2],delta(r[2],r[1]),r[3],r[4],delta(r[4],r[3]),r[17]]),[14,18,18,18,18,18,20,42]);
title(dc,'新旧 GT 重复 ID 对比','旧版读取所选模板目录的历史报告；空白表示没有对应旧 well，差值不补零。');
const total = 6+d.compare.length;
totals(dc,total,['B','C','D','E','F','G'],6,d.compare.some(r=>r[1]===null)?['B','D','E','G']:[]);
const dm = d.books.dup['口径说明'].slice(1).filter(r=>r[0]).map(r=>r.slice(0,2));
const setMethod = (rows,name,text) => {const r=rows.find(r=>r[0]===name);if(r)r[1]=text;else rows.push([name,text]);};
setMethod(dm,'扫描范围',`递归读取 ${d.outputs} 的评测 summary.json、position_matches.csv 及 summary 指向的原始 XML，共 ${d.sources.length} 个来源。`);
setMethod(dm,'数据性质',`只读分析；生成时间 ${d.scan_time}（本机时间）。原始 XML 与 CSV 已逐对象核对；表格为固定快照。`);
setMethod(dm,'核对结果',`${d.dup_events.length} 个重复事件、${d.dup_objects.length} 个实际重复对象；汇总与明细一致。`);
setMethod(dm,'旧版来源',d.template_dir+'/GT_XML重复ID汇总.xlsx');
const ms=table(dup,'口径说明',1,['项目','说明'],dm,[24,105]);
ms.getRange(`B2:B${dm.length+1}`).format.wrapText=true;ms.getRange(`A2:B${dm.length+1}`).format.rowHeight=48;
await save(dup,'duplicates');

const w=workbook(Object.keys(d.books.disp));
const main=d.books.disp['Well汇总'], header=name=>d.books.disp[name][0];
const s=table(w,'Well汇总',5,main[4],d.summary.map(r=>[r[1],r[6],r[7],r[17],r[18],r[19],r[16],r[10],r[11],r[15],r[26]]),[13,15,16,15,15,16,14,17,17,19,14]);
title(s,'GT 单帧位移异常汇总',`${d.sources.length} 个 XML / ${d.summary.length} 个 well / ${d.summary.reduce((a,r)=>a+r[3],0)} 帧；各井阈值 = 有效连续帧位移中位数 × 4。`);
for(let r=6;r<total;r++)s.getRange(`E${r}:F${r}`).formulas=[[`=IF(B${r}=0,"",D${r}/B${r})`,`=IF(C${r}=0,"",D${r}/C${r})`]];
totals(s,total,['B','C','D','G']);
s.getRange(`E${total}:F${total}`).formulas=[[`=IF(B${total}=0,"",D${total}/B${total})`,`=IF(C${total}=0,"",D${total}/C${total})`]];
fmt(s,['E','F'],6,total,'0.0%');fmt(s,['H','I','J'],6,total-1,'0.00');
const dr=total+5,sr=dr+d.summary.length+5;
s.getRange(`A${dr-2}`).values=[['分布与排除情况']];
table(w,'Well汇总',dr,main.find(r=>r[0]==='Well'&&r[1]==='有效帧对'),d.summary.map(r=>[r[1],r[9],r[12],r[13],r[14],r[8],r[20],r[21],r[5],r[4],r[26]]),[13,15,19,21,21,17,17,20,18,19,14]);
fmt(s,['C','D','E'],dr+1,dr+d.summary.length,'0.00');
s.getRange(`A${sr-2}`).values=[['敏感性复核：整条排除重复 ID track，重新计算各井阈值']];
table(w,'Well汇总',sr,main.find(r=>r[0]==='Well'&&r[1]==='剩余可评估 track').slice(0,5),d.summary.map(r=>[r[1],r[23],r[22],r[23]?r[22]/r[23]:null,r[24]]),[13,22,19,21,21]);
fmt(s,['D'],sr+1,sr+d.summary.length,'0.0%');fmt(s,['E'],sr+1,sr+d.summary.length,'0.00');
const e=table(w,'异常跳变明细',1,header('异常跳变明细'),d.events,[13,14,12,12,17,17,14,15,18,18,18,18,17,17,15,18,15,17,18,18,19,19,43,43,22,42,13,13]);
fmt(e,['E','F','G','H','I','J','K','L','M','N','P','Q'],2,d.events.length+1,'0.00');fmt(e,['O'],2,d.events.length+1,'0.0%');
if(d.events.length)e.getRange(`G2:G${d.events.length+1}`).conditionalFormats.add('cellIs',{operator:'greaterThanOrEqual',formula:5,format:{fill:'#FCE4D6',font:{bold:true,color:'#9C2F18'}}});
const t=table(w,'Track汇总',1,header('Track汇总'),d.tracks,[13,14,21,17,19,20,20,13,13,15,17,17,17,20,22,42]);fmt(t,['E'],2,d.tracks.length+1,'0.00');
const x=table(w,'排除帧对',1,header('排除帧对'),d.exclusions,[13,14,12,12,17,24,17,17,24,24,43,43,22,42,14,14]);
for(const sh of [e,t,x])sh.freezePanes.freezeColumns(2);
const c=table(w,'与旧GT对比',5,d.books.disp['与旧GT对比'][4],d.compare.map(r=>[r[0],r[5],r[6],delta(r[6],r[5]),r[7],r[8],delta(r[8],r[7]),r[9],r[10],r[11],r[12],r[13],r[14],r[15],r[16],r[17]]),[14,17,17,17,19,19,18,21,21,18,18,21,21,26,28,42]);
title(c,'新旧 GT 位移异常对比','各版本独立估计阈值；可评估帧对及 ID 可变化，数量变化不等于逐例新增或修复。');
fmt(c,['H','I'],6,total-1,'0.0%');fmt(c,['J','K','L','M'],6,total-1,'0.00');
totals(c,total,['B','C','D','E','F','G','N','O'],6,d.compare.some(r=>r[5]===null)?['B','D','E','G','N','O']:[]);
const sourceRows=d.books.disp['口径与来源'];
const methods=sourceRows.slice(1,sourceRows.findIndex(r=>r[0]==='来源编号')).filter(r=>r[0]).map(r=>r.slice(0,2));
setMethod(methods,'范围',`输入为 ${d.outputs} 的全部评测输出。原始 XML 提供轮廓与坐标，position_matches.csv 用于逐对象版本核对；位移只统计 cell。`);
setMethod(methods,'输出与复现',`生成时间 ${d.scan_time}；固定快照，不随 XML 自动刷新。仅保留两个最终 Excel。`);
setMethod(methods,'检验','CSV 与 XML 对象集合、track_id、帧号和质心逐项一致；事件与汇总核对一致；保存后回读明细数值。');
setMethod(methods,'旧版来源',d.template_dir+'/GT_单帧位移异常分析_20260915.xlsx');
table(w,'口径与来源',1,['项目','定义'],methods,[23,105]);
const m=table(w,'口径与来源',methods.length+4,sourceRows.find(r=>r[0]==='来源编号'),d.sources,[23,42,15,110,23,68,14,19,19,20]);
m.getRange(`B2:B${methods.length+1}`).format.wrapText=true;m.getRange(`A2:B${methods.length+1}`).format.rowHeight=75;
await save(w,'displacements');
'''

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--outputs',type=Path,default=HERE/'outputs')
    parser.add_argument('--template-dir',type=Path,default=HERE/'outputs')
    parser.add_argument('--output-dir',type=Path,default=HERE/'outputs')
    parser.add_argument('--gt-root',type=Path,help='Explicitly relocate GT XMLs; CSV must still match.')
    parser.add_argument('--date',default=datetime.now().strftime('%Y%m%d'))
    parser.add_argument('--overwrite',action='store_true')
    parser.add_argument('--node',type=Path)
    parser.add_argument('--node-modules',type=Path)
    args = parser.parse_args()
    try:
        datetime.strptime(args.date,'%Y%m%d')
        books = read_books(args.template_dir)
        jobs = discover(args.outputs,args.gt_root)
        data,pairs = analyze(jobs)
        data.update(books=books,compare=compare(data,pairs,books),scan_time=datetime.now().isoformat(timespec='seconds'),
                    outputs=str(args.outputs.resolve()),template_dir=str(args.template_dir.resolve()))
        export(data,args)
        print(f"Runs={len(jobs)}, duplicate events={len(data['dup_events'])}, jumps={len(data['events'])}, anomalous tracks={sum(r[17] for r in data['summary'])}")
    except (OSError,ValueError,KeyError,subprocess.CalledProcessError) as exc:
        parser.exit(2,f'GT report error: {exc}\n')


if __name__ == '__main__':
    main()
